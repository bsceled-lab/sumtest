import streamlit as st
from docx import Document
import re
import random
import hashlib

st.set_page_config(
    page_title="MCQ Master - Online Quiz",
    page_icon="📝",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown("""
<style>
.main-title {
    font-size: 2.6rem;
    font-weight: 800;
    margin-bottom: 0.2rem;
}
.subtitle {
    color: #666;
    margin-bottom: 1.2rem;
}
.question-card {
    padding: 1rem 1.2rem;
    border: 1px solid #ddd;
    border-radius: 12px;
    margin-bottom: 0.8rem;
}
.correct {
    color: #16803c;
    font-weight: 700;
}
.wrong {
    color: #c62828;
    font-weight: 700;
}
</style>
""", unsafe_allow_html=True)


def clean_text(text):
    """Normalize Word text while preserving useful punctuation."""
    if text is None:
        return ""
    text = str(text).replace("\xa0", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def strip_duplicate_question_number(text):
    """Remove duplicated numbering such as '1. 1. Question text'."""
    text = clean_text(text)
    pattern = r"^(?:(?:Q(?:uestion)?\s*)?\d+\s*[\.\):\-]?\s*){2}"
    return re.sub(pattern, "", text, count=1, flags=re.IGNORECASE).strip()


def is_question_start(text):
    text = clean_text(text)
    patterns = [
        r"^\d+\s*[\.\)]\s+",
        r"^Q(?:uestion)?\s*\d+\s*[\.\):\-]\s*",
        r"^Q\.\s*\d+\s*[\.\):\-]?\s*",
    ]
    return any(re.match(p, text, re.IGNORECASE) for p in patterns)


def normalize_question(text):
    text = strip_duplicate_question_number(text)
    text = re.sub(
        r"^(?:Q(?:uestion)?\s*)?\d+\s*[\.\):\-]\s*",
        "",
        text,
        flags=re.IGNORECASE,
    )
    text = re.sub(r"^Q\.\s*\d+\s*[\.\):\-]?\s*", "", text, flags=re.IGNORECASE)
    return clean_text(text)


def parse_option(text):
    """Return (letter, option text) for A-D option lines."""
    text = clean_text(text)
    match = re.match(r"^([A-Da-d])\s*[\.\):\-]\s*(.+)$", text)
    if match:
        return match.group(1).upper(), clean_text(match.group(2))

    # Also accept 'A Option text' when the source omits punctuation.
    match = re.match(r"^([A-Da-d])\s+(.+)$", text)
    if match:
        return match.group(1).upper(), clean_text(match.group(2))

    return None, None


def parse_answer(text):
    text = clean_text(text)
    match = re.match(
        r"^(?:Answer|Correct\s*Answer|Correct)\s*[:=\-]?\s*([A-Da-d])\b",
        text,
        re.IGNORECASE,
    )
    if match:
        return match.group(1).upper()

    # Accept a standalone answer letter.
    match = re.fullmatch(r"([A-Da-d])", text)
    if match:
        return match.group(1).upper()

    return None


def extract_lines_from_docx(uploaded_file):
    """Read paragraphs and table cells from a DOCX file."""
    document = Document(uploaded_file)
    lines = []

    for paragraph in document.paragraphs:
        value = clean_text(paragraph.text)
        if value:
            lines.append(value)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                value = clean_text(cell.text)
                if value:
                    lines.append(value)

    return lines


def parse_docx(uploaded_file):
    """
    Parse MCQs from a DOCX containing question, A-D options and Answer lines.
    Returns a list of dictionaries.
    """
    lines = extract_lines_from_docx(uploaded_file)
    questions = []
    current = None

    def save_current():
        nonlocal current
        if not current:
            return

        question_text = clean_text(current.get("question", ""))
        options = current.get("options", {})
        answer = current.get("answer")

        if question_text and all(letter in options for letter in "ABCD") and answer in "ABCD":
            questions.append({
                "question": question_text,
                "options": [options["A"], options["B"], options["C"], options["D"]],
                "answer": "ABCD".index(answer),
            })
        current = None

    for line in lines:
        answer = parse_answer(line)
        option_letter, option_text = parse_option(line)

        if is_question_start(line):
            save_current()
            current = {
                "question": normalize_question(line),
                "options": {},
                "answer": None,
            }
            continue

        if current is None:
            continue

        if option_letter:
            current["options"][option_letter] = option_text
            continue

        if answer:
            current["answer"] = answer
            continue

        # Some Word files wrap a question over multiple lines.
        # Do not append arbitrary lines after options/answer.
        if not current["options"] and current["question"]:
            current["question"] = clean_text(
                current["question"] + " " + line
            )

    save_current()

    # Remove exact duplicates while preserving order.
    unique = []
    seen = set()
    for item in questions:
        key = (
            item["question"].lower(),
            tuple(option.lower() for option in item["options"]),
            item["answer"],
        )
        if key not in seen:
            seen.add(key)
            unique.append(item)

    return unique


def reset_quiz():
    st.session_state.questions = []
    st.session_state.answers = {}
    st.session_state.submitted = False
    st.session_state.file_id = None
    st.session_state.quiz_order = []


if "questions" not in st.session_state:
    reset_quiz()

st.markdown('<div class="main-title">📝 MCQ Master</div>', unsafe_allow_html=True)
st.markdown(
    '<div class="subtitle">Upload a Word document containing MCQs, take the quiz, and view your score.</div>',
    unsafe_allow_html=True,
)

with st.sidebar:
    st.header("Quiz Settings")
    uploaded_file = st.file_uploader(
        "Upload MCQ Word file",
        type=["docx"],
        help="The document should contain questions, A-D options, and Answer lines.",
    )
    shuffle_questions = st.checkbox("Shuffle questions", value=False)

if uploaded_file is not None:
    file_bytes = uploaded_file.getvalue()
    file_id = hashlib.md5(file_bytes).hexdigest()

    if st.session_state.file_id != file_id:
        try:
            questions = parse_docx(uploaded_file)
        except Exception as exc:
            reset_quiz()
            st.error("The Word file could not be read.")
            st.exception(exc)
            st.stop()

        st.session_state.questions = questions
        st.session_state.file_id = file_id
        st.session_state.answers = {}
        st.session_state.submitted = False

        order = list(range(len(questions)))
        if shuffle_questions:
            random.shuffle(order)
        st.session_state.quiz_order = order

    elif len(st.session_state.quiz_order) != len(st.session_state.questions):
        order = list(range(len(st.session_state.questions)))
        if shuffle_questions:
            random.shuffle(order)
        st.session_state.quiz_order = order

    questions = st.session_state.questions

    if not questions:
        st.warning(
            "No complete MCQs were found. Each question must have A, B, C, D options and an Answer line."
        )
        st.stop()

    st.success(f"{len(questions)} MCQs loaded successfully.")

    if not st.session_state.submitted:
        st.subheader("Quiz")

        with st.form("mcq_quiz_form"):
            for display_no, question_index in enumerate(
                st.session_state.quiz_order, start=1
            ):
                item = questions[question_index]

                st.markdown(
                    f'<div class="question-card"><strong>Q{display_no}. {item["question"]}</strong></div>',
                    unsafe_allow_html=True,
                )

                selected = st.radio(
                    "Select one answer:",
                    options=item["options"],
                    index=None,
                    key=f"q_{question_index}",
                    label_visibility="collapsed",
                )
                st.session_state.answers[question_index] = selected

            submitted = st.form_submit_button(
                "✅ Submit Quiz",
                use_container_width=True,
                type="primary",
            )

        if submitted:
            st.session_state.submitted = True
            st.rerun()

    else:
        correct = 0
        wrong = 0
        unanswered = 0

        for index, item in enumerate(questions):
            selected = st.session_state.answers.get(index)
            correct_answer = item["options"][item["answer"]]

            if selected is None:
                unanswered += 1
            elif selected == correct_answer:
                correct += 1
            else:
                wrong += 1

        total = len(questions)
        percentage = (correct / total * 100) if total else 0

        st.subheader("📊 Quiz Result")

        col1, col2, col3, col4 = st.columns(4)
        col1.metric("Total", total)
        col2.metric("Correct", correct)
        col3.metric("Wrong", wrong)
        col4.metric("Unanswered", unanswered)

        st.progress(percentage / 100)
        st.write(f"**Score: {correct}/{total} ({percentage:.1f}%)**")

        if percentage >= 80:
            st.success("Excellent performance! 🎉")
        elif percentage >= 50:
            st.info("Good effort. Keep practicing. 👍")
        else:
            st.warning("Keep practicing and review the answers. 📚")

        st.subheader("🔎 Answer Review")

        for display_no, question_index in enumerate(
            st.session_state.quiz_order, start=1
        ):
            item = questions[question_index]
            selected = st.session_state.answers.get(question_index)
            correct_answer = item["options"][item["answer"]]

            st.markdown(f"**Q{display_no}. {item['question']}**")

            if selected is None:
                st.write("Your answer: **Not answered**")
                st.write(f"Correct answer: **{correct_answer}**")
            elif selected == correct_answer:
                st.markdown(
                    f'<span class="correct">✓ Correct: {selected}</span>',
                    unsafe_allow_html=True,
                )
            else:
                st.markdown(
                    f'<span class="wrong">✗ Your answer: {selected}</span>',
                    unsafe_allow_html=True,
                )
                st.write(f"Correct answer: **{correct_answer}**")

            st.divider()

        if st.button("🔄 Retake Quiz", use_container_width=True):
            st.session_state.answers = {}
            st.session_state.submitted = False

            order = list(range(len(questions)))
            if shuffle_questions:
                random.shuffle(order)
            st.session_state.quiz_order = order

            st.rerun()

else:
    st.info("👈 Upload a .docx MCQ file from the sidebar to begin.")
